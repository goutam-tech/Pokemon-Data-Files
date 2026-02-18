"""
Generate Pokemon.xlsx for PokemonReviewApp Database
Matches the Pokemon.cs model structure
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime, timedelta
import random
import os

def create_pokemon_file():
    wb = Workbook()
    ws = wb.active
    ws.title = "Pokemon"
    
    headers = ["Id", "Name", "BirthDate"]
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(1, col, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")
    
    base_date = datetime(2020, 1, 1)
    
    pokemon_names = [
        "Bulbasaur", "Ivysaur", "Venusaur", "Charmander", "Charmeleon",
        "Charizard", "Squirtle", "Wartortle", "Blastoise", "Caterpie",
        "Metapod", "Butterfree", "Weedle", "Kakuna", "Beedrill",
        "Pidgey", "Pidgeotto", "Pidgeot", "Rattata", "Raticate",
        "Spearow", "Fearow", "Ekans", "Arbok", "Pikachu",
        "Raichu", "Sandshrew", "Sandslash", "Nidoran♀", "Nidorina",
        "Nidoqueen", "Nidoran♂", "Nidorino", "Nidoking", "Clefairy",
        "Clefable", "Vulpix", "Ninetales", "Jigglypuff", "Wigglytuff",
        "Zubat", "Golbat", "Oddish", "Gloom", "Vileplume",
        "Paras", "Parasect", "Venonat", "Venomoth", "Diglett",
        "Dugtrio", "Meowth", "Persian", "Psyduck", "Golduck",
        "Mankey", "Primeape", "Growlithe", "Arcanine", "Poliwag",
        "Poliwhirl", "Poliwrath", "Abra", "Kadabra", "Alakazam",
        "Machop", "Machoke", "Machamp", "Bellsprout", "Weepinbell",
        "Victreebel", "Tentacool", "Tentacruel", "Geodude", "Graveler",
        "Golem", "Ponyta", "Rapidash", "Slowpoke", "Slowbro",
        "Magnemite", "Magneton", "Farfetch'd", "Doduo", "Dodrio",
        "Seel", "Dewgong", "Grimer", "Muk", "Shellder",
        "Cloyster", "Gastly", "Haunter", "Gengar", "Onix",
        "Drowzee", "Hypno", "Krabby", "Kingler", "Voltorb",
        "Electrode", "Exeggcute", "Exeggutor", "Cubone", "Marowak",
        "Hitmonlee", "Hitmonchan", "Lickitung", "Koffing", "Weezing",
        "Rhyhorn", "Rhydon", "Chansey", "Tangela", "Kangaskhan",
        "Horsea", "Seadra", "Goldeen", "Seaking", "Staryu",
        "Starmie", "Mr. Mime", "Scyther", "Jynx", "Electabuzz",
        "Magmar", "Pinsir", "Tauros", "Magikarp", "Gyarados",
        "Lapras", "Ditto", "Eevee", "Vaporeon", "Jolteon",
        "Flareon", "Porygon", "Omanyte", "Omastar", "Kabuto",
        "Kabutops", "Aerodactyl", "Snorlax", "Articuno", "Zapdos",
        "Moltres", "Dratini", "Dragonair", "Dragonite", "Mewtwo",
        "Mew", "Chikorita", "Bayleef", "Meganium", "Cyndaquil",
        "Quilava", "Typhlosion", "Totodile", "Croconaw", "Feraligatr",
        "Sentret", "Furret", "Hoothoot", "Noctowl", "Ledyba",
        "Ledian", "Spinarak", "Ariados", "Crobat", "Chinchou",
        "Lanturn", "Pichu", "Cleffa", "Igglybuff", "Togepi",
        "Togetic", "Natu", "Xatu", "Mareep", "Flaaffy",
        "Ampharos", "Bellossom", "Marill", "Azumarill", "Sudowoodo",
        "Politoed", "Hoppip", "Skiploom", "Jumpluff", "Aipom",
        "Sunkern", "Sunflora", "Yanma", "Wooper", "Quagsire",
        "Espeon", "Umbreon", "Murkrow", "Slowking", "Misdreavus",
        "Unown", "Wobbuffet", "Girafarig", "Pineco", "Forretress",
        "Dunsparce", "Gligar", "Steelix", "Snubbull", "Granbull",
        "Qwilfish", "Scizor", "Shuckle", "Heracross", "Sneasel",
        "Teddiursa", "Ursaring", "Slugma", "Magcargo", "Swinub",
        "Piloswine", "Corsola", "Remoraid", "Octillery", "Delibird",
        "Mantine", "Skarmory", "Houndour", "Houndoom", "Kingdra",
        "Phanpy", "Donphan", "Porygon2", "Stantler", "Smeargle",
        "Tyrogue", "Hitmontop", "Smoochum", "Elekid", "Magby",
        "Miltank", "Blissey", "Raikou", "Entei", "Suicune",
        "Larvitar", "Pupitar", "Tyranitar", "Lugia", "Ho-Oh",
        "Celebi"
    ]
    
    pokemon_data = []
    for i, name in enumerate(pokemon_names, 1):
        days_offset = random.randint(0, 1825)
        birth_date = (base_date + timedelta(days=days_offset)).strftime("%Y-%m-%d")
        pokemon_data.append([i, name, birth_date])
    
    for row_data in pokemon_data:
        ws.append(row_data)
    
    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 15
    
    os.makedirs('report', exist_ok=True)
    output_path = 'report/XLSX/Pokemon.xlsx'
    wb.save(output_path)
    print(f"✓ Pokemon.xlsx created with {len(pokemon_data)} Pokemon")
    return output_path

if __name__ == "__main__":
    print("=" * 60)
    print("GENERATING POKEMON DATA")
    print("=" * 60)
    create_pokemon_file()
    print("\n✓ File ready for database import!")