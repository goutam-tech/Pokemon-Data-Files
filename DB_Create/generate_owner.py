"""
Generate Owner.xlsx for PokemonReviewApp Database
Matches the Owner.cs model structure
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import os

def create_owner_file():
    wb = Workbook()
    ws = wb.active
    ws.title = "Owners"
    
    headers = ["Id", "FirstName", "LastName", "Gym", "CountryId"]
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(1, col, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")
    
    owners = [
        [1, "Ash", "Ketchum", "Pallet Town", 1],
        [2, "Misty", "Williams", "Cerulean City Gym", 1],
        [3, "Brock", "Harrison", "Pewter City Gym", 1],
        [4, "Gary", "Oak", "Pallet Town", 1],
        [5, "Professor", "Oak", "Pallet Town Lab", 1],
        [6, "Lt.", "Surge", "Vermilion City Gym", 2],
        [7, "Erika", "Celadon", "Celadon City Gym", 1],
        [8, "Sabrina", "Natsume", "Saffron City Gym", 1],
        [9, "Koga", "Fuchsia", "Fuchsia City Gym", 1],
        [10, "Blaine", "Cinnabar", "Cinnabar Island Gym", 2],
        [11, "Giovanni", "Rocket", "Viridian City Gym", 6],
        [12, "Lance", "Dragon", "Indigo Plateau", 1],
        [13, "Lorelei", "Prima", "Elite Four", 1],
        [14, "Bruno", "Shiba", "Elite Four", 1],
        [15, "Agatha", "Kikuko", "Elite Four", 1],
        [16, "Falkner", "Violet", "Violet City Gym", 1],
        [17, "Bugsy", "Azalea", "Azalea Town Gym", 1],
        [18, "Whitney", "Goldenrod", "Goldenrod City Gym", 1],
        [19, "Morty", "Ecruteak", "Ecruteak City Gym", 1],
        [20, "Chuck", "Cianwood", "Cianwood City Gym", 1],
        [21, "Jasmine", "Olivine", "Olivine City Gym", 1],
        [22, "Pryce", "Mahogany", "Mahogany Town Gym", 1],
        [23, "Clair", "Blackthorn", "Blackthorn City Gym", 1],
        [24, "Will", "Elite", "Elite Four Johto", 1],
        [25, "Karen", "Dark", "Elite Four Johto", 1],
        [26, "May", "Maple", "Littleroot Town", 1],
        [27, "Dawn", "Berlitz", "Twinleaf Town", 1],
        [28, "Serena", "Yvonne", "Vaniville Town", 4],
        [29, "Clemont", "Lumiose", "Lumiose City Gym", 4],
        [30, "Cynthia", "Shirona", "Sinnoh Champion", 1],
        [31, "Steven", "Stone", "Hoenn Champion", 1],
        [32, "Wallace", "Sootopolis", "Sootopolis City Gym", 1],
        [33, "Roxanne", "Rustboro", "Rustboro City Gym", 1],
        [34, "Brawly", "Dewford", "Dewford Town Gym", 1],
        [35, "Wattson", "Mauville", "Mauville City Gym", 1],
        [36, "Flannery", "Lavaridge", "Lavaridge Town Gym", 1],
        [37, "Norman", "Petalburg", "Petalburg City Gym", 1],
        [38, "Winona", "Fortree", "Fortree City Gym", 1],
        [39, "Tate", "Mossdeep", "Mossdeep City Gym", 1],
        [40, "Liza", "Mossdeep", "Mossdeep City Gym", 1],
        [41, "Juan", "Sootopolis", "Sootopolis City Gym", 7],
        [42, "Roark", "Oreburgh", "Oreburgh City Gym", 1],
        [43, "Gardenia", "Eterna", "Eterna City Gym", 1],
        [44, "Maylene", "Veilstone", "Veilstone City Gym", 1],
        [45, "Crasher", "Wake", "Pastoria City Gym", 1],
        [46, "Fantina", "Hearthome", "Hearthome City Gym", 4],
        [47, "Byron", "Canalave", "Canalave City Gym", 1],
        [48, "Candice", "Snowpoint", "Snowpoint City Gym", 1],
        [49, "Volkner", "Sunyshore", "Sunyshore City Gym", 1],
        [50, "Paul", "Veilstone", "Trainer", 1],
    ]
    
    for row_data in owners:
        ws.append(row_data)
    
    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 25
    ws.column_dimensions['E'].width = 12
    
    os.makedirs('report', exist_ok=True) 
    output_path = 'report/Owner.xlsx'
    wb.save(output_path)
    print(f"✓ Owner.xlsx created with {len(owners)} owners")
    return output_path

if __name__ == "__main__":
    print("=" * 60)
    print("GENERATING OWNER DATA")
    print("=" * 60)
    create_owner_file()
    print("\n✓ File ready for database import!")