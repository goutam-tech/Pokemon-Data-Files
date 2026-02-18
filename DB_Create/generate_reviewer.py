"""
Generate Reviewer.xlsx for PokemonReviewApp Database
Matches the Reviewer.cs model structure
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import os

def create_reviewer_file():
    wb = Workbook()
    ws = wb.active
    ws.title = "Reviewers"
    
    headers = ["Id", "FirstName", "LastName"]
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(1, col, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="C55A11", end_color="C55A11", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")
    
    reviewers = [
        [1, "Ash", "Ketchum"],
        [2, "Misty", "Williams"],
        [3, "Brock", "Harrison"],
        [4, "Professor", "Oak"],
        [5, "Gary", "Oak"],
        [6, "Lance", "Dragon"],
        [7, "Cynthia", "Shirona"],
        [8, "Steven", "Stone"],
        [9, "May", "Maple"],
        [10, "Dawn", "Berlitz"],
        [11, "Serena", "Yvonne"],
        [12, "Clemont", "Lumiose"],
        [13, "Paul", "Veilstone"],
        [14, "Tracey", "Sketchit"],
        [15, "Max", "Maple"],
        [16, "Iris", "Village"],
        [17, "Cilan", "Striaton"],
        [18, "N", "Harmonia"],
        [19, "Lillie", "Aether"],
        [20, "Gladion", "Aether"],
        [21, "Goh", "Vermilion"],
        [22, "Chloe", "Cerise"],
        [23, "Professor", "Kukui"],
        [24, "Professor", "Elm"],
        [25, "Professor", "Birch"],
        [26, "Professor", "Rowan"],
        [27, "Professor", "Juniper"],
        [28, "Professor", "Sycamore"],
        [29, "Nurse", "Joy"],
        [30, "Officer", "Jenny"],
    ]
    
    for row_data in reviewers:
        ws.append(row_data)
    
    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 20
    
    os.makedirs('report', exist_ok=True)
    output_path = 'report/XLSX/Reviewer.xlsx'
    wb.save(output_path)
    print(f"✓ Reviewer.xlsx created with {len(reviewers)} reviewers")
    return output_path

if __name__ == "__main__":
    print("=" * 60)
    print("GENERATING REVIEWER DATA")
    print("=" * 60)
    create_reviewer_file()
    print("\n✓ File ready for database import!")