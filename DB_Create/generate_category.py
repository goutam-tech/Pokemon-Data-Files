"""
Generate Category.xlsx for PokemonReviewApp Database
Matches the Category.cs model structure
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import os

def create_category_file():
    wb = Workbook()
    ws = wb.active
    ws.title = "Categories"
    
    headers = ["Id", "Name"]
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(1, col, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")
    
    categories = [
        [1, "Normal"],
        [2, "Fire"],
        [3, "Water"],
        [4, "Electric"],
        [5, "Grass"],
        [6, "Ice"],
        [7, "Fighting"],
        [8, "Poison"],
        [9, "Ground"],
        [10, "Flying"],
        [11, "Psychic"],
        [12, "Bug"],
        [13, "Rock"],
        [14, "Ghost"],
        [15, "Dragon"],
        [16, "Dark"],
        [17, "Steel"],
        [18, "Fairy"],
    ]
    
    for row_data in categories:
        ws.append(row_data)
    
    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 20
    
    os.makedirs('report', exist_ok=True) 
    output_path = 'report/Category.xlsx'
    wb.save(output_path)
    print(f"✓ Category.xlsx created with {len(categories)} categories")
    return output_path

if __name__ == "__main__":
    print("=" * 60)
    print("GENERATING CATEGORY DATA")
    print("=" * 60)
    create_category_file()
    print("\n✓ File ready for database import!")