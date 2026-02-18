"""
Generate PokemonOwner.xlsx for PokemonReviewApp Database
Matches the PokemonOwner.cs model structure (Many-to-Many junction table)
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import os

def create_pokemon_owner_file():
    wb = Workbook()
    ws = wb.active
    ws.title = "PokemonOwners"
    
    headers = ["PokemonId", "OwnerId"]
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(1, col, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="305496", end_color="305496", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")
    
    pokemon_owners = [
        [25, 1],
        [6, 1],
        [1, 1],
        [7, 1],
        [18, 1],
        [12, 1],
        [143, 1],
        [146, 1],
        [155, 1],
        [160, 1],

        [120, 2],
        [121, 2],
        [54, 2],
        [116, 2],
        [131, 2],
        [9, 2],

        [95, 3],
        [74, 3],
        [76, 3],
        [42, 3],
        [208, 3],

        [3, 4],
        [9, 4],
        [133, 4],
        [65, 4],
        [59, 4],
        [149, 4],

        [133, 5],
        [151, 5],
        [150, 5],

        [26, 6],
        [125, 6],
        [82, 6],
        [101, 6],

        [3, 7],
        [45, 7],
        [71, 7],
        [114, 7],

        [65, 8],
        [97, 8],
        [122, 8],
        [103, 8],

        [110, 9],
        [89, 9],
        [73, 9],
        [49, 9],

        [59, 10],
        [126, 10],
        [78, 10],
        [136, 10],

        [31, 11],
        [34, 11],
        [112, 11],
        [150, 11],

        [149, 12],
        [148, 12],
        [130, 12],
        [142, 12],
        [6, 12],

        [87, 13],
        [91, 13],
        [144, 13],
        [131, 13],
        [124, 13],

        [68, 14],
        [95, 14],
        [57, 14],
        [106, 14],
        [107, 14],

        [94, 15],
        [42, 15],
        [93, 15],
        [24, 15],

        [157, 26],
        [9, 26],
        [123, 26],
        [127, 26],

        [68, 27],
        [154, 27],

        [136, 28],
        [196, 28],

        [149, 30],
        [3, 30],

        [212, 31],
        [208, 31],

        [16, 1],
        [19, 4],
        [129, 46],
        [50, 3],
        [52, 16],
        [100, 6],
        [35, 7],
        [39, 13],
        [132, 48],
        [137, 50],
    ]

    
    for row_data in pokemon_owners:
        ws.append(row_data)
    
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 12
    
    os.makedirs('report', exist_ok=True)
    output_path = 'report/XLSX/PokemonOwner.xlsx'
    wb.save(output_path)
    print(f"✓ PokemonOwner.xlsx created with {len(pokemon_owners)} mappings")
    return output_path

if __name__ == "__main__":
    print("=" * 60)
    print("GENERATING POKEMON-OWNER JUNCTION TABLE")
    print("=" * 60)
    create_pokemon_owner_file()
    print("\n✓ File ready for database import!")