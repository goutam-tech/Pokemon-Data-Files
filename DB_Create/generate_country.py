"""
Generate Country.xlsx for PokemonReviewApp Database
Matches the Country.cs model structure
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import os

def create_country_file():
    wb = Workbook()
    ws = wb.active
    ws.title = "Countries"
    
    headers = ["Id", "Name"]
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(1, col, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")
    
    countries = [
        [1, "Japan"],
        [2, "United States"],
        [3, "United Kingdom"],
        [4, "France"],
        [5, "Germany"],
        [6, "Italy"],
        [7, "Spain"],
        [8, "Canada"],
        [9, "Australia"],
        [10, "Brazil"],
        [11, "South Korea"],
        [12, "China"],
        [13, "India"],
        [14, "Mexico"],
        [15, "Netherlands"],
    ]
    
    for row_data in countries:
        ws.append(row_data)
    
    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 25
    
    os.makedirs('report', exist_ok=True)
    output_path = 'report/Country.xlsx'
    wb.save(output_path)
    print(f"✓ Country.xlsx created with {len(countries)} countries")
    return output_path

if __name__ == "__main__":
    print("=" * 60)
    print("GENERATING COUNTRY DATA")
    print("=" * 60)
    create_country_file()
    print("\n✓ File ready for database import!")